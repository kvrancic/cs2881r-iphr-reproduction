"""Beats -> .pptx. One slide per beat; every object is a picture named !!id so Morph matches it across slides.
Morph within a scene, a fade between scenes. No fonts are needed on the presenting machine: all text is pictures."""
import io, hashlib
from lxml import etree
from pptx import Presentation
from pptx.util import Emu
from pptx.dml.color import RGBColor
from PIL import Image
from pal import C
from render import placement

EMU = 31750            # one native pixel (384 across 13.333 in) is exactly 31750 EMU
K = 10                 # export scale for 384-grid art: 3840x2160 equivalent, so 1080p and 4K both get whole pixels

NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}

FADE = """<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med"><p:fade/></p:transition>"""

MORPH = """<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
 xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
 xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main">
 <mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">
  <p:transition spd="slow" p14:dur="{ms}"><p159:morph option="byObject"/></p:transition>
 </mc:Choice>
 <mc:Fallback><p:transition spd="fast"><p:fade/></p:transition></mc:Fallback>
</mc:AlternateContent>"""


def _png(img, f, cache={}):
    h = hashlib.sha1(img.tobytes() + bytes(f"{img.size}{f}", "ascii")).hexdigest()
    if h not in cache:
        big = img if f == 1 else img.resize((img.width * f, img.height * f), Image.NEAREST)
        b = io.BytesIO(); big.save(b, "PNG", optimize=True)
        cache[h] = b.getvalue()
    return io.BytesIO(cache[h])


def _set_alpha(pic, alpha):
    if alpha >= 0.999:
        return
    blip = pic._element.find(".//a:blip", NS)
    a = etree.SubElement(blip, "{%s}alphaModFix" % NS["a"])
    a.set("amt", str(int(alpha * 100000)))


def notes_text(b, nxt):
    return (b.say + "\n\n" if b.say else "") + ("NEXT: " + nxt)


def build(beats, out, title="IPHR reproduction"):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(384 * EMU), Emu(216 * EMU)
    layout = prs.slide_layouts[6]
    for i, b in enumerate(beats):
        s = prs.slides.add_slide(layout)
        fill = s.background.fill; fill.solid(); fill.fore_color.rgb = RGBColor(*C[b.bg])
        for o in sorted(b.objs.values(), key=lambda o: o.z):
            if o.alpha <= 0:
                continue
            x, y, w, h = placement(o, b.cam)
            f = max(1, round(K * o.scale))
            pic = s.shapes.add_picture(_png(o.img, f), Emu(round(x * EMU)), Emu(round(y * EMU)),
                                       Emu(round(w * EMU)), Emu(round(h * EMU)))
            pic.name = "!!" + o.id
            _set_alpha(pic, o.alpha)
        nxt = next((x.say for x in beats[i + 1:] if x.say), "(end)")
        s.notes_slide.notes_text_frame.text = notes_text(b, nxt)
        el = s._element
        for t in el.findall("p:timing", NS):
            el.remove(t)
        fade = i == 0 or b.fade
        t = etree.fromstring(FADE if fade else MORPH.format(ms=int(b.morph * 1000)))
        clr = el.find("p:clrMapOvr", NS)
        (clr.addnext(t) if clr is not None else el.find("p:cSld", NS).addnext(t))
    sz = prs.part._element.find("p:sldSz", NS)
    if sz is not None and "type" in sz.attrib:
        del sz.attrib["type"]            # the slide size is custom 16:9, not "On-screen Show (4:3)"
    prs.core_properties.title = title
    prs.save(out)
    return out, len(beats)
