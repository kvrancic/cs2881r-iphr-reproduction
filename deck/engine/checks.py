"""Checks run on every build: the speaker notes are exactly the transcript, no text collides with anything,
every slide has a transition, and Morph names line up."""
import os
import numpy as np
from px import W, H

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SCRIPT = os.path.join(ROOT, "talk", "transcript.md")


def spoken():
    """The words said, from talk/transcript.md: every paragraph that is not a heading."""
    out = []
    for para in open(SCRIPT).read().split("\n\n"):
        p = para.strip()
        if not p or p.startswith("#"):
            continue
        out.append(p)
    return " ".join(" ".join(out).split())


def notes_match(beats):
    said = " ".join(" ".join(b.say for b in beats if b.say).split())
    want = spoken()
    empty = [i + 1 for i, b in enumerate(beats) if not b.say.strip()]
    if said == want and not empty:
        return True, f"speaker notes match talk/transcript.md word for word ({len(want.split())} words, {len(beats)} slides)"
    if empty:
        return False, f"slides without spoken words: {empty}"
    i = next((k for k, (a, b) in enumerate(zip(said, want)) if a != b), min(len(said), len(want)))
    return False, f"notes differ from the transcript at char {i}:\n  deck:       ...{said[max(0, i - 60):i + 60]}\n  transcript: ...{want[max(0, i - 60):i + 60]}"


def morph_report(beats):
    """For each transition: ids that carry over, appear and disappear."""
    lines = []
    for i in range(1, len(beats)):
        a, b = beats[i - 1], beats[i]
        A, B = set(a.objs) - {"ground"}, set(b.objs) - {"ground"}
        kind = "fade" if b.fade else "morph"
        lines.append(f"{i:2d}->{i + 1:2d} {kind:5s} keep {len(A & B):2d}  in {sorted(B - A)}  out {sorted(A - B)}")
    return lines


# ------------------------------------------------------------------------------------------------ layout
def _mask(o, cam):
    """The object's opaque pixels on the 384x216 screen, at 2x for sub-pixel chart objects."""
    from render import placement
    from PIL import Image
    k = 2
    x, y, w, h = placement(o, cam)
    a = o.img.split()[3].resize((max(1, round(w * k)), max(1, round(h * k))), Image.NEAREST)
    m = np.zeros((H * k, W * k), bool)
    X, Y = round(x * k), round(y * k)
    arr = np.asarray(a) > 0
    x0, y0 = max(0, X), max(0, Y)
    x1, y1 = min(W * k, X + arr.shape[1]), min(H * k, Y + arr.shape[0])
    if x1 > x0 and y1 > y0:
        m[y0:y1, x0:x1] = arr[y0 - Y:y1 - Y, x0 - X:x1 - X]
    off = arr.sum() - m.sum()
    return m, off


def check_layout(beats):
    """Text (the cue, every t_* label and every chart text element) must not touch any other object, except the
    backdrops it is meant to sit on, and must stay on screen. Returns a list of issues."""
    from story.common import BACKDROPS, BACKDROP_PREFIX
    import charts
    textish = lambda i: i == "cue" or i.startswith("t_") or (i.startswith("ch_") and i[3:] in charts.TEXTS)
    backdrop = lambda i: i in BACKDROPS or i.startswith(BACKDROP_PREFIX) or (i.startswith("ch_") and i[3:] in charts.PANELS | charts.GRIDS | charts.PLATES)
    issues = [f"chart element runs past the 960x540 grid: {x}" for x in charts.OVERFLOW]
    for n, b in enumerate(beats, 1):
        objs = {i: o for i, o in b.objs.items() if o.alpha > 0.05}
        masks = {i: _mask(o, b.cam) for i, o in objs.items() if i != "ground"}
        for i, (m, off) in masks.items():
            if not textish(i):
                continue
            if off > 0:
                issues.append(f"slide {n}: {i} runs off screen ({off} px)")
            for j, (m2, _) in masks.items():
                if j == i or backdrop(j) or (textish(j) and j < i):
                    continue
                ov = int((m & m2).sum())
                if ov > 2:
                    issues.append(f"slide {n}: {i} touches {j} ({ov} px at 2x)")
    return issues


# ------------------------------------------------------------------------------------------------ the .pptx
def validate_pptx(path, beats):
    """Round trip through python-pptx: slide count, 16:9, names, a transition on every slide, notes."""
    from pptx import Presentation
    from lxml import etree
    from build_pptx import notes_text
    prs = Presentation(path)
    errs = []
    if len(prs.slides) != len(beats):
        errs.append(f"{len(prs.slides)} slides, expected {len(beats)}")
    if abs(prs.slide_width / prs.slide_height - 16 / 9) > 1e-3:
        errs.append("not 16:9")
    n_morph = n_fade = 0
    for i, (s, b) in enumerate(zip(prs.slides, beats), 1):
        names = [sh.name for sh in s.shapes]
        if any(not n.startswith("!!") for n in names):
            errs.append(f"slide {i}: unnamed picture")
        if len(set(names)) != len(names):
            errs.append(f"slide {i}: duplicate names")
        if any(sh.shape_type != 13 for sh in s.shapes):
            errs.append(f"slide {i}: an object that is not a picture")
        xml = etree.tostring(s._element).decode()
        if "p159:morph" in xml:
            n_morph += 1
        elif "<p:fade" in xml:
            n_fade += 1
        else:
            errs.append(f"slide {i}: no transition")
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        nxt = next((x.say for x in beats[i:] if x.say), "(end)")
        if notes != notes_text(b, nxt):
            errs.append(f"slide {i}: notes are not the spoken words + NEXT")
    return errs, f"{len(prs.slides)} slides, {n_morph} Morph, {n_fade} fade, every slide has a transition"


# ------------------------------------------------------------------------------------------------ claims
def claims():
    """Every number and checkable statement the speaker makes, checked against the result files. Returns [(ok, text)]."""
    import json, os
    import charts as ch
    c, R, N, G = ch.CHECKS, ch.RUNS, ch.NOISE, ch.GEO["conditions"]
    CF = ch.J("cited_facts.json")
    f2 = c["figure2"]
    out = []
    n = ch.N_PAIRS
    out.append((4500 <= n < 5000 and n == 4834, f"'almost five thousand pairs' / 'the data hold 4,834': {n:,} pairs"))
    out.append((len(ch.FIG2) == 16, f"'sixteen models': {len(ch.FIG2)} bars"))
    g = f2["openai/gpt-4o-mini"]["paper_pct"]
    out.append((12.5 <= g < 14.5 and g == max(v["paper_pct"] for v in f2.values()),
                f"'GPT-4o mini is the highest at about thirteen percent': {g}%"))
    th = [f2[k]["paper_pct"] for k, _, t in ch.FIG2 if t]
    out.append((sum(v < 0.5 for v in th) * 2 > len(th), f"'most thinking models below half a percent': {sum(v < 0.5 for v in th)} of {len(th)}"))
    rec = [k for k, v in f2.items() if v["source"] == "recomputed"]
    same = sum(f2[k]["paper_count"] == f2[k]["our_count"] for k in rec)
    out.append((len(rec) == 8 and same == 8, f"'eight models ... exactly their numbers': {same} of {len(rec)} recomputed counts match"))
    out.append((len(rec) == 8 and same == 8, f"'for the eight models with released data, Figure 2's counts recompute exactly': {same} of {len(rec)}"))
    bad = [k for k, v in f2.items() if round(100 * v["paper_count"] / ch.PLOT_DIVISOR, 2) != v["paper_pct"]]
    out.append((not bad, f"'divides by 4,892 pairs': count / 4,892 gives every plotted bar ({len(f2) - len(bad)} of {len(f2)})"))
    m = c["retest_mean_of_models"]
    out.append((abs(m["thr04"] - 0.765) < abs(m["thr05"] - 0.765),
                f"'the retest uses a looser cutoff': 0.4 keeps {m['thr04']:.1%}, nearer the paper's 76.5% than 0.5 ({m['thr05']:.1%})"))
    rs = N["rank_stability"]
    out.append((rs["settings_with_any_swap"] > 0, f"'its own table shows that it does [change]': swaps in {rs['settings_with_any_swap']} of {rs['settings'] - 1} other settings"))
    out.append((0.18 <= R["base"]["openai/gpt-4o-mini"]["n_pairs"] / n < 0.23, f"'a fifth of the pairs': {R['base']['openai/gpt-4o-mini']['n_pairs']:,} of {n:,}"))
    p, nj = ch.judge_pct()
    out.append((p == 99, f"'about ninety-nine percent of their answers': {p}% over {nj} answers"))
    lg = {lab: N["paired"][k]["lean_gap"] for k, lab in ch.FOUR}
    out.append((all(abs(v["ours"] / v["paper"] - 1) < 0.25 for v in lg.values()),
                "'leans about as strongly': " + ", ".join(f"{l} {v['paper']:.3f}/{v['ours']:.3f}" for l, v in lg.items())))
    sm = {lab: (R["paper_same_pairs"][k]["n_unf"], R["base"][k]["n_unf"]) for k, lab in ch.FOUR[:2]}
    out.append((all(abs(b / a - 1) < 0.3 for a, b in sm.values()), f"'GPT-4o mini and Llama about the same number of flags': {sm}"))
    d = N["paired"][ch.DS]
    ds = (d["paper_run_vs_ours"]["flags"], d["paper_answers_our_labels_vs_ours"]["flags"], R["base"][ch.DS]["n_unf"])
    out.append((ds[0] > ds[1] > ds[2], f"'DeepSeek gets fewer, part of that our labeller': paper {ds[0]}, paper's answers with our labels {ds[1]}, ours {ds[2]}"))
    gx = N["paired"]["openai/gpt-4o-mini"]["one_sided_lean_recurrence"]
    out.append((gx["mean_in_other_run"] > 2 * gx["null_975"],
                f"'GPT-4o mini ... also did better on that order in ours, beyond its general lean': "
                f"{gx['mean_in_other_run']:.2f} against noise up to {gx['null_975']:.2f}"))
    out.append((gx["anti_mean_in_other_run"] > 2 * gx["anti_null_975"] and gx["anti_mean_in_other_run"] >= gx["mean_in_other_run"],
                f"'it did so just as often on pairs that went against its lean': against the lean "
                f"{gx['anti_mean_in_other_run']:.2f} (noise up to {gx['anti_null_975']:.2f}), with the lean {gx['mean_in_other_run']:.2f}"))
    rk = N["rate_by_kind_paper_full"]
    gm = rk["openai/gpt-4o-mini"]
    out.append((0.18 <= gm["geography"]["rate"] < 0.22 and 0.005 <= gm["other"]["rate"] < 0.02,
                f"'GPT-4o mini a fifth of geography, one percent of the rest': {gm['geography']['rate']:.1%}, {gm['other']['rate']:.1%}"))
    rev = [k for k in ("qwen/qwq-32b", "meta-llama/Llama-3.3-70B-Instruct", "meta-llama/Llama-3.1-70B") if rk[k]["geography"]["rate"] < rk[k]["other"]["rate"]]
    out.append((len(rev) == 3, f"'QwQ and Llama the other way round': {len(rev)} of 3 higher on the other questions"))
    rg, ro = N["ranking_geography"], N["ranking_other"]
    out.append((rg[0] in ro[len(ro) // 2:] and rg[-1] in ro[:len(ro) // 2],
                f"'the top and bottom of the ranking swap': top on geography {rg[0]} is in the bottom half on the rest; bottom {rg[-1]} in the top half"))
    lt = c["lean_by_kind_total"]
    out.append((lt["geography_NO"] > 5 * lt["geography_YES"] and lt["other_YES"] > 5 * lt["other_NO"],
                f"'geography leans no, dates/sizes/ages lean yes': {lt}"))
    a, b = ch.nongeo_total()
    out.append((a == 57 and b == 0, f"'from fifty-seven to zero across four models': {a} -> {b} over {len(ch.FACT4)} models"))
    lr = ch.J("recall_probe.json")["meta-llama/llama-3.3-70b-instruct"]   # checks the spoken figure only; not shown on screen
    out.append((0.65 <= lr["known_order_flagged"] < 0.75, f"'Llama ordered about seven in ten of its flagged pairs correctly': {lr['known_order_flagged']:.0%} of {lr['flagged']}"))
    out.append((0.8 <= lr["known_order_unflagged"] < 0.9, f"'against eight and a half in ten for its other pairs': {lr['known_order_unflagged']:.1%} of {lr['unflagged']}"))
    lc = CF["meta-llama/llama-3.3-70b-instruct | other"]
    out.append((lc["wrong_states_wrong_order"] > 0.5, f"'in its wrong answers the numbers it writes support the wrong answer': {lc['wrong_states_wrong_order']:.0%} of {lc['wrong_answers']}"))
    q = ch.GEO["distances"]["paper | openai/gpt-4o-mini"]["flag_rate_by_asked_quartile"]
    out.append((q[-1] == max(q), f"'flags most common for places far apart': by distance quarter {', '.join(f'{v:.1%}' for v in q)}"))
    out.append((G["base"]["flags"] == 131, f"'with the paper's wording it has 131 flags': {G['base']['flags']} of {G['base']['pairs']}"))
    out.append((135 <= G["paraphrase"]["flags"] <= 145, f"'left the flags at about 140': {G['paraphrase']['flags']} (paper's wording {G['base']['flags']})"))
    out.append((G["clarified"]["flags"] == 31, f"'cut them to 31': {G['clarified']['flags']}"))
    out.append((G["strict"]["flags"] == 667 and G["strict"]["pairs"] == 684, f"'raised them to 667 of 684': {G['strict']['flags']} of {G['strict']['pairs']}"))
    import re
    paper = open(os.path.join(ROOT, "scratch", "paper.txt")).read()
    m_ = re.search(r"GPT-4o-mini\s+(\d+\.\d+)%", paper[paper.index("near-identical longitudes"):])
    ab = float(m_.group(1)) if m_ else -1
    out.append((round(ab) == ch.ABLATION_PCT == 9, f"'the paper's own test ... still flags GPT-4o mini on nine percent': {ab}% in the paper's ablation"))
    out.append((lc["wrong_says_unsure"] > 0.5, f"'Llama often says it is unsure and answers anyway': {lc['wrong_says_unsure']:.0%} of its wrong answers"))
    return out
